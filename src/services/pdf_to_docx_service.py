
"""
Project : Document Converter
Project ID : 034

PDF to DOCX Conversion Service
"""

import shutil
import tempfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from src.config import OUTPUT_DIR


@dataclass(frozen=True)
class PDFToDOCXResult:
    success: bool
    input_path: str
    output_path: str
    message: str
    error_details: str = ""


class PDFToDOCXService:
    """Extract selectable PDF text into a Word document."""

    def convert_pdf_to_docx(
        self,
        input_path: str | Path,
        output_path: str | Path | None = None,
    ) -> PDFToDOCXResult:
        """Convert a PDF file into a DOCX document."""

        source = Path(input_path).expanduser()
        destination = (
            Path(output_path).expanduser()
            if output_path is not None
            else OUTPUT_DIR / f"{source.stem}.docx"
        )

        temporary_path: Path | None = None

        try:
            source = source.resolve()
            destination = destination.resolve()

            validation_error = self._validate_paths(
                source, destination
            )
            if validation_error:
                return self._failure(
                    source, destination, validation_error
                )

            destination.parent.mkdir(
                parents=True, exist_ok=True
            )

            reader = PdfReader(str(source))

            if reader.is_encrypted:
                return self._failure(
                    source,
                    destination,
                    "Encrypted PDFs are not supported.",
                )

            if len(reader.pages) == 0:
                return self._failure(
                    source,
                    destination,
                    "The PDF contains no pages.",
                )

            extracted_pages: list[str] = []

            for page_number, page in enumerate(
                reader.pages, start=1
            ):
                try:
                    page_text = page.extract_text()
                except Exception as exc:
                    return self._failure(
                        source,
                        destination,
                        f"Text extraction failed on page "
                        f"{page_number}: {exc}",
                    )

                extracted_pages.append(page_text or "")

            if not any(text.strip() for text in extracted_pages):
                return self._failure(
                    source,
                    destination,
                    "No selectable text was found. "
                    "Scanned PDFs requiring OCR are not supported.",
                )

            word_document = Document()

            for page_index, page_text in enumerate(
                extracted_pages
            ):
                for line in page_text.splitlines():
                    if line.strip():
                        word_document.add_paragraph(line)
                    else:
                        word_document.add_paragraph()

                if page_index < len(extracted_pages) - 1:
                    word_document.add_page_break()

            with tempfile.NamedTemporaryFile(
                suffix=".docx",
                prefix="pdf_to_docx_",
                dir=str(destination.parent),
                delete=False,
            ) as temporary_file:
                temporary_path = Path(temporary_file.name)

            word_document.save(str(temporary_path))

            if destination.exists():
                return self._failure(
                    source,
                    destination,
                    "The output file already exists. "
                    "It was not overwritten.",
                )

            shutil.move(
                str(temporary_path),
                str(destination),
            )
            temporary_path = None

            if not destination.is_file():
                return self._failure(
                    source,
                    destination,
                    "The output DOCX file was not created.",
                )

            if destination.stat().st_size == 0:
                destination.unlink(missing_ok=True)
                return self._failure(
                    source,
                    destination,
                    "The output DOCX file is empty.",
                )

            return PDFToDOCXResult(
                success=True,
                input_path=str(source),
                output_path=str(destination),
                message=(
                    "PDF to DOCX conversion completed "
                    "successfully. Extracted text was saved."
                ),
            )

        except Exception as exc:
            return self._failure(
                source,
                destination,
                f"{type(exc).__name__}: {exc}",
            )

        finally:
            if (
                temporary_path is not None
                and temporary_path.exists()
            ):
                temporary_path.unlink(missing_ok=True)

    @staticmethod
    def _validate_paths(
        source: Path,
        destination: Path,
    ) -> str | None:
        """Validate input and output paths."""

        if not source.exists():
            return "The input PDF file does not exist."

        if not source.is_file():
            return "The input path is not a file."

        if source.suffix.lower() != ".pdf":
            return "The input file must have a .pdf extension."

        if destination.suffix.lower() != ".docx":
            return "The output file must have a .docx extension."

        if source == destination:
            return "Input and output paths must be different."

        if destination.exists():
            return (
                "The output file already exists. "
                "It will not be overwritten."
            )

        return None

    @staticmethod
    def _failure(
        source: Path,
        destination: Path,
        message: str,
    ) -> PDFToDOCXResult:
        """Build a consistent conversion failure result."""

        return PDFToDOCXResult(
            success=False,
            input_path=str(source),
            output_path=str(destination),
            message="PDF to DOCX conversion failed.",
            error_details=message,
        )


# Public interface
# Class: PDFToDOCXService
# Method: convert_pdf_to_docx(input_path, output_path=None)
# Result: PDFToDOCXResult
# Dependencies: pypdf, python-docx, standard library,
#               src.config.OUTPUT_DIR
