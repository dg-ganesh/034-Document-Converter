
"""
Project : Document Converter
Project ID : 034

Text and Word Document Conversion Service
"""

from dataclasses import dataclass
from pathlib import Path

from docx import Document

from src.config import DEFAULT_ENCODING, OUTPUT_DIR


@dataclass(frozen=True)
class TextConversionResult:
    success: bool
    input_path: str
    output_path: str
    message: str
    error_details: str = ""


class TextConversionService:
    """Convert between DOCX and plain-text files."""

    def convert_docx_to_txt(
        self,
        input_path: str | Path,
        output_path: str | Path | None = None,
    ) -> TextConversionResult:
        """Extract paragraphs from a DOCX file into UTF-8 text."""

        source = Path(input_path).expanduser()
        destination = (
            Path(output_path).expanduser()
            if output_path is not None
            else OUTPUT_DIR / f"{source.stem}.txt"
        )

        try:
            source = source.resolve()
            destination = destination.resolve()

            error = self._validate_paths(
                source, destination, ".docx", ".txt"
            )
            if error:
                return self._failure(
                    source, destination, error
                )

            document = Document(str(source))
            lines = [paragraph.text for paragraph in document.paragraphs]
            text_content = "\n".join(lines)

            destination.parent.mkdir(
                parents=True, exist_ok=True
            )
            destination.write_text(
                text_content,
                encoding=DEFAULT_ENCODING,
            )

            return self._success(
                source,
                destination,
                "DOCX to TXT conversion completed successfully.",
            )

        except Exception as exc:
            return self._failure(
                source,
                destination,
                f"{type(exc).__name__}: {exc}",
            )

    def convert_txt_to_docx(
        self,
        input_path: str | Path,
        output_path: str | Path | None = None,
    ) -> TextConversionResult:
        """Convert UTF-8 text lines into DOCX paragraphs."""

        source = Path(input_path).expanduser()
        destination = (
            Path(output_path).expanduser()
            if output_path is not None
            else OUTPUT_DIR / f"{source.stem}.docx"
        )

        try:
            source = source.resolve()
            destination = destination.resolve()

            error = self._validate_paths(
                source, destination, ".txt", ".docx"
            )
            if error:
                return self._failure(
                    source, destination, error
                )

            text_content = source.read_text(
                encoding=DEFAULT_ENCODING
            )

            document = Document()

            for line in text_content.splitlines():
                document.add_paragraph(line)

            destination.parent.mkdir(
                parents=True, exist_ok=True
            )
            document.save(str(destination))

            return self._success(
                source,
                destination,
                "TXT to DOCX conversion completed successfully.",
            )

        except Exception as exc:
            return self._failure(
                source,
                destination,
                f"{type(exc).__name__}: {exc}",
            )

    @staticmethod
    def _validate_paths(
        source: Path,
        destination: Path,
        expected_input_extension: str,
        expected_output_extension: str,
    ) -> str | None:
        """Validate conversion input and output paths."""

        if not source.exists():
            return "The input file does not exist."

        if not source.is_file():
            return "The input path is not a file."

        if source.suffix.lower() != expected_input_extension:
            return (
                f"Input file must have the "
                f"{expected_input_extension} extension."
            )

        if destination.suffix.lower() != expected_output_extension:
            return (
                f"Output file must have the "
                f"{expected_output_extension} extension."
            )

        if source == destination:
            return "Input and output paths must be different."

        if destination.exists():
            return (
                "The output file already exists. "
                "It will not be overwritten."
            )

        return None

    @staticmethod
    def _success(
        source: Path,
        destination: Path,
        message: str,
    ) -> TextConversionResult:
        """Build a consistent successful result."""

        return TextConversionResult(
            success=True,
            input_path=str(source),
            output_path=str(destination),
            message=message,
        )

    @staticmethod
    def _failure(
        source: Path,
        destination: Path,
        error_details: str,
    ) -> TextConversionResult:
        """Build a consistent failure result."""

        return TextConversionResult(
            success=False,
            input_path=str(source),
            output_path=str(destination),
            message="Text conversion failed.",
            error_details=error_details,
        )


# Public interface
# Class: TextConversionService
# Methods:
#   convert_docx_to_txt(input_path, output_path=None)
#   convert_txt_to_docx(input_path, output_path=None)
# Result: TextConversionResult
# Dependencies: python-docx, standard library, src.config
